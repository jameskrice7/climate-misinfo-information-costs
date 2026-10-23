"""Create the journal's editable highlights using python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor

ROOT=Path(__file__).resolve().parent
highlights=[
 'Sector-specific climate content connects political information to energy finance.',
 'Greater exposure is associated with wider spreads under local market-day controls.',
 'The association is imprecise after sector-month conditions are absorbed.',
 'A sector-sensitive index preserves differences across fossil and renewable firms.',
 'Timing and alternative-outcome checks do not establish additional robust effects.'
]
assert all(len(h)<=85 for h in highlights)
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.space_after=Pt(12)
title=doc.styles['Title'];title.font.name='Times New Roman';title.font.size=Pt(16);title.font.color.rgb=RGBColor(0,0,0)
doc.add_paragraph('Highlights',style='Title')
doc.add_paragraph('The Information Costs of Climate Misinformation')
doc.add_paragraph('Evidence from Global Financial Data')
for h in highlights:doc.add_paragraph(h,style='List Bullet')
# Remove inherited template rules as well as direct paragraph borders.
for root in [doc._element,doc.styles._element]:
    for element in list(root.iter()):
        if element.tag.endswith('}pBdr'):
            element.getparent().remove(element)
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.title='Highlights'
doc.save(ROOT/'Highlights.docx')
(ROOT/'Highlights.txt').write_text('\n'.join('- '+h for h in highlights)+'\n',encoding='utf-8')
print('Highlight lengths:',[len(h) for h in highlights])
